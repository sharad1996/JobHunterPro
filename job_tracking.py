"""
Excel (xlsx) job tracking + skip rules for already-actioned roles.

Columns are looked up by header name so you can reorder columns in Excel.
Mark \"Applied on portal\", \"Application email sent\", or isEmailed=TRUE to skip application emails.
Mark \"Follow-up email sent\" or isFollowed=TRUE to skip follow-up emails.

Each search upserts rows by Job URL (with HR email when found). Fill HR email in the sheet and run
`import_hr_from_xlsx_to_db` via `python3 main.py --sync-from-xlsx` to push addresses into SQLite for sending.
"""

import os
from datetime import datetime
from urllib.parse import urlparse, urlunparse

import config

try:
    from openpyxl import Workbook, load_workbook
except ImportError:  # pragma: no cover
    Workbook = load_workbook = None  # type: ignore


HEADERS = [
    "Job URL",
    "Company",
    "Job title",
    "Platform",
    "Region",
    "Company website",
    "Headcount note",
    "Application email sent",
    "Follow-up email sent",
    "isEmailed",
    "isFollowed",
    "Applied on portal",
    "HR email",
    "Notes",
    "Updated at",
]


def normalize_job_url(url: str) -> str:
    if not url or not isinstance(url, str):
        return ""
    u = url.strip()
    try:
        p = urlparse(u)
        host = (p.netloc or "").lower()
        if host.startswith("www."):
            host = host[4:]
        path = (p.path or "").rstrip("/") or "/"
        q = ""  # strip query for stable matching
        return urlunparse((p.scheme.lower(), host, path, "", q, ""))
    except Exception:
        return u.lower().rstrip("/")


def _xlsx_path() -> str:
    return getattr(config, "JOB_TRACKING_XLSX", "job_tracking.xlsx")


def _ensure_openpyxl():
    if Workbook is None:
        raise RuntimeError("openpyxl is required for Excel tracking. Run: pip install openpyxl")


def ensure_workbook(path: str = None) -> str:
    _ensure_openpyxl()
    path = path or _xlsx_path()
    if os.path.exists(path):
        return path
    parent = os.path.dirname(os.path.abspath(path))
    if parent:
        os.makedirs(parent, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Jobs"
    ws.append(HEADERS)
    wb.save(path)
    return path


def _header_map(ws) -> dict:
    row1 = [str(c.value or "").strip() for c in ws[1]]
    return {name: idx + 1 for idx, name in enumerate(row1)}


def _cell(ws, col_map: dict, row: int, header: str, default=""):
    c = col_map.get(header)
    if not c:
        return default
    v = ws.cell(row=row, column=c).value
    return "" if v is None else str(v).strip()


def _is_yes(val: str) -> bool:
    v = (val or "").strip().lower()
    return v in ("y", "yes", "true", "1", "x")


def load_skip_url_set(path: str = None) -> set:
    """URLs that should not be searched/applied again."""
    path = path or _xlsx_path()
    if not os.path.exists(path):
        return set()
    _ensure_openpyxl()
    wb = load_workbook(path, read_only=True)
    ws = wb.active
    it = ws.iter_rows(values_only=True)
    try:
        header = next(it)
    except StopIteration:
        wb.close()
        return set()
    headers = [str(h or "").strip() for h in header]
    col_map = {h: i for i, h in enumerate(headers)}
    url_i = col_map.get("Job URL")
    app_i = col_map.get("Application email sent")
    portal_i = col_map.get("Applied on portal")
    is_emailed_i = col_map.get("isEmailed")
    if url_i is None:
        wb.close()
        return set()
    skip = set()
    for row in it:
        if not row or url_i >= len(row) or not row[url_i]:
            continue
        url = normalize_job_url(str(row[url_i]))
        if not url:
            continue
        app = _is_yes(str(row[app_i]) if app_i is not None and app_i < len(row) else "")
        portal = _is_yes(str(row[portal_i]) if portal_i is not None and portal_i < len(row) else "")
        is_emailed = _is_yes(
            str(row[is_emailed_i]) if is_emailed_i is not None and is_emailed_i < len(row) else ""
        )
        if app or portal or is_emailed:
            skip.add(url)
    wb.close()
    return skip


def load_followup_skip_urls(path: str = None) -> set:
    """URLs that must not receive a follow-up (Follow-up email sent / isFollowed on the sheet)."""
    path = path or _xlsx_path()
    if not os.path.exists(path):
        return set()
    _ensure_openpyxl()
    wb = load_workbook(path, read_only=True)
    ws = wb.active
    it = ws.iter_rows(values_only=True)
    try:
        header = next(it)
    except StopIteration:
        wb.close()
        return set()
    headers = [str(h or "").strip() for h in header]
    col_map = {h: i for i, h in enumerate(headers)}
    url_i = col_map.get("Job URL")
    fu_i = col_map.get("Follow-up email sent")
    is_followed_i = col_map.get("isFollowed")
    if url_i is None:
        wb.close()
        return set()
    skip = set()
    for row in it:
        if not row or url_i >= len(row) or not row[url_i]:
            continue
        url = normalize_job_url(str(row[url_i]))
        if not url:
            continue
        fu = _is_yes(str(row[fu_i]) if fu_i is not None and fu_i < len(row) else "")
        is_followed = _is_yes(
            str(row[is_followed_i]) if is_followed_i is not None and is_followed_i < len(row) else ""
        )
        if fu or is_followed:
            skip.add(url)
    wb.close()
    return skip


def upsert_row(
    job: dict,
    job_title: str,
    path: str = None,
    *,
    application_email_sent: str = "",
    followup_sent: str = "",
    applied_portal: str = "",
    hr_email: str = "",
    notes: str = "",
) -> None:
    """Insert or update one row keyed by Job URL."""
    path = ensure_workbook(path)
    _ensure_openpyxl()
    wb = load_workbook(path)
    ws = wb.active
    cmap = _header_map(ws)
    for h in HEADERS:
        if h not in cmap:
            # append missing columns
            col = ws.max_column + 1
            ws.cell(row=1, column=col, value=h)
            cmap = _header_map(ws)

    raw_url = (job.get("url") or "").strip()
    url = normalize_job_url(raw_url)
    if not url:
        wb.close()
        return

    target_row = None
    for r in range(2, ws.max_row + 1):
        if normalize_job_url(_cell(ws, cmap, r, "Job URL")) == url:
            target_row = r
            break

    if target_row is None:
        target_row = ws.max_row + 1

    def put(header, value, overwrite=True):
        c = cmap.get(header)
        if not c:
            return
        cur = ws.cell(row=target_row, column=c).value
        if not overwrite and cur not in (None, ""):
            return
        ws.cell(row=target_row, column=c, value=value)

    put("Job URL", raw_url or job.get("url") or "", True)
    put("Company", job.get("company") or "", True)
    put("Job title", job.get("title") or job_title, True)
    put("Platform", job.get("platform") or "", True)
    put("Region", job.get("search_country") or "", True)
    website = (job.get("company_website") or job.get("domain") or "").strip()
    if website:
        put("Company website", website, False)
    hn = job.get("headcount_note") or job.get("company_size") or ""
    put("Headcount note", hn, True)
    if application_email_sent:
        put("Application email sent", application_email_sent, True)
        put("isEmailed", "TRUE", True)
    if followup_sent:
        put("Follow-up email sent", followup_sent, True)
        put("isFollowed", "TRUE", True)
    if applied_portal:
        put("Applied on portal", applied_portal, True)
    resolved_hr = (hr_email or (job.get("hr_email") or "")).strip()
    if resolved_hr:
        put("HR email", resolved_hr, True)
    if notes:
        put("Notes", notes, False)
    elif not resolved_hr:
        hint = (
            "Find HR from job URL; add email in this column or Notes, then run: "
            "python3 main.py --sync-from-xlsx"
        )
        put("Notes", hint, False)
    put("Updated at", datetime.now().strftime("%Y-%m-%d %H:%M"), True)

    wb.save(path)
    wb.close()


def bulk_upsert_initial(jobs: list, job_title: str, path: str = None) -> tuple:
    """
    Write every scraped listing to the sheet BEFORE HR-email enrichment.
    Opens the workbook once. Inserts new rows (keyed by Job URL); for existing rows,
    only fills empty cells (won't overwrite manual edits or sent flags).
    Returns (inserted, updated).
    """
    path = ensure_workbook(path)
    _ensure_openpyxl()
    wb = load_workbook(path)
    ws = wb.active
    cmap = _header_map(ws)
    for h in HEADERS:
        if h not in cmap:
            col = ws.max_column + 1
            ws.cell(row=1, column=col, value=h)
    cmap = _header_map(ws)

    existing = {}
    for r in range(2, ws.max_row + 1):
        nu = normalize_job_url(_cell(ws, cmap, r, "Job URL"))
        if nu:
            existing[nu] = r

    inserted = 0
    updated = 0
    now = datetime.now().strftime("%Y-%m-%d %H:%M")

    for job in jobs:
        raw_url = (job.get("url") or "").strip()
        nu = normalize_job_url(raw_url)
        if not nu:
            continue
        row = existing.get(nu)
        if row is None:
            row = ws.max_row + 1 if ws.max_row >= 1 else 2
            ws.cell(row=row, column=cmap["Job URL"], value=raw_url)
            existing[nu] = row
            inserted += 1
        else:
            updated += 1

        def fill(header, value, overwrite=False):
            c = cmap.get(header)
            if not c:
                return
            cur = ws.cell(row=row, column=c).value
            if not value:
                return
            if not overwrite and cur not in (None, ""):
                return
            ws.cell(row=row, column=c, value=value)

        fill("Company", job.get("company") or "")
        fill("Job title", job.get("title") or job_title)
        fill("Platform", job.get("platform") or "")
        fill("Region", job.get("search_country") or "")
        website = (job.get("company_website") or job.get("domain") or "").strip()
        if website:
            fill("Company website", website)
        hn = job.get("headcount_note") or job.get("company_size") or ""
        if hn:
            fill("Headcount note", hn)
        fill(
            "Notes",
            "Find HR from job URL; type email in HR email column, then run: "
            "python3 main.py --sync-from-xlsx",
        )
        ws.cell(row=row, column=cmap["Updated at"], value=now)

    wb.save(path)
    wb.close()
    return inserted, updated


def sync_from_database(db, path: str = None) -> None:
    """Rebuild sheet rows from SQLite (one row per job)."""
    path = path or _xlsx_path()
    parent = os.path.dirname(os.path.abspath(path))
    if parent:
        os.makedirs(parent, exist_ok=True)
    path = ensure_workbook(path)
    _ensure_openpyxl()
    jobs = db.get_all_jobs(5000)
    wb = Workbook()
    ws = wb.active
    ws.title = "Jobs"
    ws.append(HEADERS)
    for j in jobs:
        url = j.get("job_url") or ""
        app = "Y" if (j.get("email_status") == "sent") else ""
        fu = "Y" if j.get("follow_up_sent") else ""
        is_em = "TRUE" if j.get("email_status") == "sent" else ""
        is_fol = "TRUE" if j.get("follow_up_sent") else ""
        ws.append(
            [
                url,
                j.get("company") or "",
                j.get("job_title") or "",
                j.get("platform") or "",
                j.get("search_country") or "",
                j.get("domain") or "",
                "",
                app,
                fu,
                is_em,
                is_fol,
                "",
                j.get("hr_email") or "",
                "",
                (j.get("applied_at") or "")[:19],
            ]
        )
    wb.save(path)
    wb.close()


def _read_sheet_rows(path: str = None) -> tuple:
    """
    Load all data rows from the tracking sheet.
    Returns (path, list of {header: value, _row: int}).
    """
    path = path or _xlsx_path()
    if not os.path.isfile(path):
        return path, []
    _ensure_openpyxl()
    wb = load_workbook(path, read_only=True)
    ws = wb.active
    it = ws.iter_rows(values_only=True)
    try:
        header = next(it)
    except StopIteration:
        wb.close()
        return path, []

    headers = [str(h or "").strip() for h in header]
    col_map = {h: i for i, h in enumerate(headers)}
    url_i = col_map.get("Job URL")
    rows_out = []
    row_num = 1
    for row in it:
        row_num += 1
        if not row:
            continue
        url = str(row[url_i]).strip() if url_i is not None and url_i < len(row) and row[url_i] else ""
        if not url:
            continue
        entry = {"_row": row_num, "Job URL": url}
        for h in HEADERS:
            if h == "Job URL":
                continue
            i = col_map.get(h)
            if i is None or i >= len(row):
                entry[h] = ""
            else:
                entry[h] = "" if row[i] is None else str(row[i]).strip()
        rows_out.append(entry)
    wb.close()
    return path, rows_out


def _xlsx_row_to_db_payload(xr: dict) -> dict:
    return {
        "job_url": xr.get("Job URL") or "",
        "company": xr.get("Company") or "",
        "job_title": xr.get("Job title") or "",
        "platform": xr.get("Platform") or "",
        "search_country": xr.get("Region") or "",
        "company_domain": xr.get("Company website") or "",
        "hr_email": xr.get("HR email") or "",
        "application_sent": _is_yes(xr.get("Application email sent", ""))
        or _is_yes(xr.get("isEmailed", "")),
        "followup_sent": _is_yes(xr.get("Follow-up email sent", ""))
        or _is_yes(xr.get("isFollowed", "")),
    }


def _merge_db_job_into_sheet_row(ws, cmap: dict, row_num: int, db_job: dict, xlsx_row: dict):
    """Update sheet cells from DB when the sheet cell is empty or DB has sent status."""
    now = datetime.now().strftime("%Y-%m-%d %H:%M")

    def put(header, value, overwrite=False):
        c = cmap.get(header)
        if not c:
            return
        cur = ws.cell(row=row_num, column=c).value
        if not overwrite and cur not in (None, ""):
            return
        ws.cell(row=row_num, column=c, value=value)

    xlsx_row = xlsx_row or {}
    if not (xlsx_row.get("HR email") or "").strip() and (db_job.get("hr_email") or "").strip():
        put("HR email", db_job["hr_email"], True)

    if db_job.get("email_status") == "sent":
        put("Application email sent", "Y", True)
        put("isEmailed", "TRUE", True)
    if db_job.get("follow_up_sent"):
        put("Follow-up email sent", "Y", True)
        put("isFollowed", "TRUE", True)

    if not (xlsx_row.get("Company website") or "").strip() and (db_job.get("company_domain") or "").strip():
        put("Company website", db_job["company_domain"], True)
    if not (xlsx_row.get("Company") or "").strip():
        put("Company", db_job.get("company") or "", True)
    if not (xlsx_row.get("Job title") or "").strip():
        put("Job title", db_job.get("job_title") or "", True)
    if not (xlsx_row.get("Platform") or "").strip():
        put("Platform", db_job.get("platform") or "", True)
    if not (xlsx_row.get("Region") or "").strip():
        put("Region", db_job.get("search_country") or "", True)

    put("Updated at", now, True)


def sync_with_database(db, path: str = None) -> dict:
    """
    Two-way sync between job_tracking.xlsx and SQLite, keyed by normalized Job URL.
    - Sheet rows missing in DB → inserted into DB
    - DB rows missing in sheet → appended to sheet
    - Rows in both → HR email and sent flags merged (sheet HR wins if filled)
    """
    path, sheet_rows = _read_sheet_rows(path)
    path = ensure_workbook(path)
    _ensure_openpyxl()

    xlsx_by_url = {}
    for xr in sheet_rows:
        nu = normalize_job_url(xr.get("Job URL") or "")
        if nu:
            xlsx_by_url[nu] = xr

    stats = {
        "db_inserted": 0,
        "db_updated": 0,
        "db_skipped": 0,
        "xlsx_appended": 0,
        "xlsx_updated": 0,
        "sheet_rows": len(sheet_rows),
    }

    for nu, xr in xlsx_by_url.items():
        action = db.upsert_from_sheet_row(_xlsx_row_to_db_payload(xr))
        if action == "inserted":
            stats["db_inserted"] += 1
        elif action == "updated":
            stats["db_updated"] += 1
        else:
            stats["db_skipped"] += 1

    db_jobs = db.get_all_jobs(10000)
    db_by_url = {}
    for j in db_jobs:
        nu = normalize_job_url(j.get("job_url") or "")
        if nu:
            db_by_url[nu] = j

    wb = load_workbook(path)
    ws = wb.active
    cmap = _header_map(ws)
    for h in HEADERS:
        if h not in cmap:
            col = ws.max_column + 1
            ws.cell(row=1, column=col, value=h)
            cmap = _header_map(ws)

    for nu, dj in db_by_url.items():
        xr = xlsx_by_url.get(nu)
        if xr is None:
            job = {
                "url": dj.get("job_url") or "",
                "company": dj.get("company") or "",
                "title": dj.get("job_title") or "",
                "platform": dj.get("platform") or "",
                "search_country": dj.get("search_country") or "",
                "domain": dj.get("company_domain") or "",
                "hr_email": dj.get("hr_email") or "",
            }
            upsert_row(
                job,
                dj.get("job_title") or "",
                application_email_sent="Y" if dj.get("email_status") == "sent" else "",
                followup_sent="Y" if dj.get("follow_up_sent") else "",
                hr_email=dj.get("hr_email") or "",
            )
            stats["xlsx_appended"] += 1
        else:
            _merge_db_job_into_sheet_row(ws, cmap, xr["_row"], dj, xr)
            stats["xlsx_updated"] += 1

    wb.save(path)
    wb.close()
    return stats


def import_hr_from_xlsx_to_db(db, path: str = None) -> tuple:
    """
    Read Job URL + HR email from the tracking sheet and push into SQLite pending rows
    (same normalized URL). Skips rows without URL or without HR email; does not change sent rows.
    Returns (rows_updated, rows_skipped_no_match).
    """
    path = path or _xlsx_path()
    if not os.path.isfile(path):
        return 0, 0
    _ensure_openpyxl()
    wb = load_workbook(path, read_only=True)
    ws = wb.active
    it = ws.iter_rows(values_only=True)
    try:
        header = next(it)
    except StopIteration:
        wb.close()
        return 0, 0
    headers = [str(h or "").strip() for h in header]
    col_map = {h: i for i, h in enumerate(headers)}
    url_i = col_map.get("Job URL")
    hr_i = col_map.get("HR email")
    if url_i is None or hr_i is None:
        wb.close()
        return 0, 0

    updated = 0
    skipped = 0
    for row in it:
        if not row:
            continue
        url = str(row[url_i]).strip() if url_i < len(row) and row[url_i] else ""
        hr = str(row[hr_i]).strip() if hr_i < len(row) and row[hr_i] else ""
        nu = normalize_job_url(url)
        if not nu or not hr:
            continue
        n = db.update_hr_by_normalized_job_url(nu, hr, "", verified=True)
        if n:
            updated += n
        else:
            skipped += 1
    wb.close()
    return updated, skipped


def _company_cache_key(company: str) -> str:
    from email_finder import normalize_company_name

    return normalize_company_name(company).lower()


def enrich_hr_emails_in_xlsx(
    path: str = None,
    *,
    only_missing: bool = True,
    include_guesses: bool = None,
    limit: int = None,
    dry_run: bool = False,
) -> dict:
    """
    For rows in the tracking sheet, scrape company contact pages and fill the HR email column.
    Reuses one lookup per company name (multiple job rows share the same result).
    Returns stats dict: processed, filled, verified, guessed, skipped_has_email, not_found.
    """
    import time

    from email_finder import find_hr_email_for_company, looks_valid_email

    path = ensure_workbook(path)
    _ensure_openpyxl()
    wb = load_workbook(path)
    ws = wb.active
    cmap = _header_map(ws)
    for h in HEADERS:
        if h not in cmap:
            col = ws.max_column + 1
            ws.cell(row=1, column=col, value=h)
            cmap = _header_map(ws)

    if include_guesses is None:
        include_guesses = bool(getattr(config, "FILL_XLSX_INCLUDE_GUESSES", False))

    rows = []
    for r in range(2, ws.max_row + 1):
        company = _cell(ws, cmap, r, "Company")
        if not company:
            continue
        hr = _cell(ws, cmap, r, "HR email")
        if only_missing and hr:
            continue
        rows.append(
            (
                r,
                company,
                _cell(ws, cmap, r, "Job URL"),
                _cell(ws, cmap, r, "Company website"),
            )
        )

    if limit is not None and limit > 0:
        rows = rows[:limit]

    stats = {
        "processed": 0,
        "filled": 0,
        "verified": 0,
        "guessed": 0,
        "skipped_has_email": 0,
        "not_found": 0,
        "companies_scraped": 0,
    }

    if only_missing:
        for r in range(2, ws.max_row + 1):
            if _cell(ws, cmap, r, "Company") and _cell(ws, cmap, r, "HR email"):
                stats["skipped_has_email"] += 1

    if not rows:
        wb.close()
        return stats

    print(f"\n📧 Filling HR emails in {path} ({len(rows)} row(s) to process)…")
    company_cache: dict = {}
    now = datetime.now().strftime("%Y-%m-%d %H:%M")
    total = len(rows)

    for i, (row_num, company, job_url, website) in enumerate(rows, 1):
        stats["processed"] += 1
        cache_key = _company_cache_key(company)
        if cache_key not in company_cache:
            print(f"  [{i}/{total}] {company[:56]}", end="  ", flush=True)
            try:
                result = find_hr_email_for_company(
                    company,
                    job_url,
                    website,
                    allow_guesses=include_guesses,
                )
            except Exception as exc:
                print(f"→ error ({exc})")
                result = {}
            else:
                em = (result.get("hr_email") or "").strip()
                if em and looks_valid_email(em):
                    tag = "contact page" if result.get("hr_email_verified") else "guessed"
                    print(f"→ {em}  [{tag}]")
                else:
                    print("→ Not found")
            company_cache[cache_key] = result
            stats["companies_scraped"] += 1
            time.sleep(0.25)
        else:
            result = company_cache[cache_key]

        em = (result.get("hr_email") or "").strip()
        if not em or not looks_valid_email(em):
            stats["not_found"] += 1
            continue
        if not result.get("hr_email_verified") and not include_guesses:
            stats["not_found"] += 1
            continue

        stats["filled"] += 1
        if result.get("hr_email_verified"):
            stats["verified"] += 1
        else:
            stats["guessed"] += 1

        if dry_run:
            continue

        def put(header, value, overwrite=True):
            c = cmap.get(header)
            if not c:
                return
            cur = ws.cell(row=row_num, column=c).value
            if not overwrite and cur not in (None, ""):
                return
            ws.cell(row=row_num, column=c, value=value)

        put("HR email", em, True)
        dom = (result.get("domain") or "").strip()
        if dom:
            put("Company website", dom, False)
        src = (result.get("contact_source_url") or "").strip()
        note = f"HR from contact page ({src})" if src else "HR from company contact page"
        if not result.get("hr_email_verified"):
            note = f"Guessed HR email ({em}); verify before sending"
        put("Notes", note, True)
        put("Updated at", now, True)

    if not dry_run and stats["filled"]:
        wb.save(path)
    wb.close()
    return stats


def filter_jobs_not_skipped(jobs: list, skip_urls: set) -> tuple:
    """Remove jobs whose URL is in skip_urls."""
    if not skip_urls:
        return jobs, 0
    out = []
    n = 0
    for j in jobs:
        u = normalize_job_url(j.get("url") or "")
        if u and u in skip_urls:
            n += 1
            continue
        out.append(j)
    if n:
        print(f"  ⏭️  Skipped {n} job(s) already marked applied / emailed in {_xlsx_path()}.")
    return out, n
